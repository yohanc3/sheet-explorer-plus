from __future__ import annotations

import csv
import io
import json
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
import unicodedata
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

import requests
from flask import Flask, jsonify, render_template, request
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
DATABASE = Path(os.environ.get("SHEET_EXPLORER_DB", DATA_DIR / "sheet_explorer.db"))
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_NOTEBOOK_BYTES = 15 * 1024 * 1024
FILE_ID_RE = re.compile(r"^[A-Za-z0-9_-]{10,200}$")
MASTER_SHEET_EXPORT_URL = os.environ.get(
    "MASTER_SHEET_EXPORT_URL",
    "https://docs.google.com/spreadsheets/d/"
    "1cWTdj4B2X45MApBdnaGFF2VjLPXhM7opJSuZ6Zon3PE/export?format=csv&gid=2007683856",
)
SUBMISSION_COLUMNS = (
    "timestamp", "title", "first_name", "last_name", "full_name", "time", "difficulty",
    "confident", "needswork", "suggestions", "corrections", "locals", "share",
)

app = Flask(__name__)
app.config.update(
    MAX_CONTENT_LENGTH=MAX_UPLOAD_BYTES,
    MASTER_SHEET_EXPORT_URL=MASTER_SHEET_EXPORT_URL,
)


@contextmanager
def db() -> Iterator[sqlite3.Connection]:
    DATABASE.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def init_db() -> None:
    with db() as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS classes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE COLLATE NOCASE,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                class_id INTEGER NOT NULL REFERENCES classes(id) ON DELETE CASCADE,
                name TEXT NOT NULL COLLATE NOCASE,
                position INTEGER NOT NULL DEFAULT 0,
                UNIQUE(class_id, name)
            );
            CREATE INDEX IF NOT EXISTS idx_students_class_position
                ON students(class_id, position);
            CREATE TABLE IF NOT EXISTS submissions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT,
                title TEXT NOT NULL DEFAULT '',
                first_name TEXT NOT NULL DEFAULT '',
                last_name TEXT NOT NULL DEFAULT '',
                full_name TEXT NOT NULL DEFAULT '',
                time TEXT NOT NULL DEFAULT '',
                difficulty TEXT NOT NULL DEFAULT '',
                confident TEXT NOT NULL DEFAULT '',
                needswork TEXT NOT NULL DEFAULT '',
                suggestions TEXT NOT NULL DEFAULT '',
                corrections TEXT NOT NULL DEFAULT '',
                locals TEXT NOT NULL DEFAULT '',
                share TEXT NOT NULL DEFAULT ''
            );
            CREATE INDEX IF NOT EXISTS idx_submissions_title_name
                ON submissions(title, full_name);
            """
        )
        connection.execute("PRAGMA optimize")


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def display_student_name(value: Any) -> str:
    """Preserve a student's spelling while removing invisible/duplicate spacing."""
    return " ".join(unicodedata.normalize("NFKC", clean(value)).split())


def student_name_key(value: Any) -> str:
    """Return the stable identity used to group a student's submissions."""
    return display_student_name(value).casefold()


def canonicalize_submission_names(rows: list[sqlite3.Row]) -> list[dict[str, Any]]:
    submissions = [dict(row) for row in rows]
    variants: dict[str, dict[str, list[int]]] = {}

    # Rows arrive newest first. Prefer a spelling with readable capitalization,
    # then frequency and recency, without rewriting names such as McDonald.
    for index, submission in enumerate(submissions):
        display_name = display_student_name(submission["full_name"])
        key = student_name_key(display_name)
        entry = variants.setdefault(key, {}).setdefault(display_name, [0, index])
        entry[0] += 1

    canonical: dict[str, str] = {}
    for key, spellings in variants.items():
        canonical[key] = max(
            spellings,
            key=lambda name: (
                sum(word[:1].isupper() for word in name.split()),
                not name.isupper(),
                spellings[name][0],
                -spellings[name][1],
            ),
        )

    for submission in submissions:
        key = student_name_key(submission["full_name"])
        submission["student_key"] = key
        submission["full_name"] = canonical.get(key, display_student_name(submission["full_name"]))
    return submissions


def timestamp_value(value: Any, epoch: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value, epoch).isoformat()
        except (TypeError, ValueError):
            pass
    text = clean(value)
    for pattern in ("%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, pattern).isoformat()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).isoformat()
    except ValueError:
        return text


def hyperlink_value(cell: Any) -> str:
    if cell.hyperlink and cell.hyperlink.target:
        return clean(cell.hyperlink.target)
    value = clean(cell.value)
    formula = re.match(r'^=HYPERLINK\(\s*"([^"]+)"', value, flags=re.IGNORECASE)
    return formula.group(1) if formula else value


def submission_record(row: dict[str, Any], share: str, epoch: Any = None) -> tuple[str, ...]:
    first_name, last_name = clean(row.get("first_name")), clean(row.get("last_name"))
    return (
        timestamp_value(row.get("timestamp"), epoch), clean(row.get("title")), first_name, last_name,
        f"{first_name} {last_name}".strip(), clean(row.get("time")), clean(row.get("difficulty")),
        clean(row.get("confident")), clean(row.get("needswork")), clean(row.get("suggestions")),
        clean(row.get("corrections")), clean(row.get("locals")), clean(share),
    )


def validate_headers(headers: list[str]) -> None:
    required = {"title", "first_name", "last_name", "share"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")


def records_from_xlsx(stream: Any) -> list[tuple[str, ...]]:
    workbook = load_workbook(stream, read_only=False, data_only=False)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows()
        raw_headers = next(rows, None)
        if not raw_headers:
            raise ValueError("The first worksheet is empty.")
        headers = [clean(cell.value).lower() for cell in raw_headers]
        validate_headers(headers)

        records = []
        for cells in rows:
            row = dict(zip(headers, (cell.value for cell in cells)))
            cell_by_header = dict(zip(headers, cells))
            if not any(clean(cell.value) for cell in cells):
                continue
            records.append(submission_record(row, hyperlink_value(cell_by_header["share"]), workbook.epoch))
        return records
    finally:
        workbook.close()


def records_from_csv(data: bytes) -> list[tuple[str, ...]]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text, newline=""))
    if not reader.fieldnames:
        raise ValueError("The Google Sheet is empty.")
    headers = [clean(header).lower() for header in reader.fieldnames]
    validate_headers(headers)
    reader.fieldnames = headers

    records = []
    for row in reader:
        if not any(clean(value) for value in row.values() if value is not None):
            continue
        records.append(submission_record(row, clean(row.get("share"))))
    if not records:
        raise ValueError("The Google Sheet does not contain any submissions.")
    return records


def replace_submissions(records: list[tuple[str, ...]]) -> None:
    with db() as connection:
        connection.execute("DELETE FROM submissions")
        connection.executemany(
            f"""INSERT INTO submissions ({', '.join(SUBMISSION_COLUMNS)})
            VALUES ({', '.join('?' for _ in SUBMISSION_COLUMNS)})""",
            records,
        )


def serialize_class(row: sqlite3.Row, connection: sqlite3.Connection) -> dict[str, Any]:
    students = connection.execute(
        "SELECT id, name, position FROM students WHERE class_id = ? ORDER BY position, name",
        (row["id"],),
    ).fetchall()
    return {"id": row["id"], "name": row["name"], "students": [dict(item) for item in students]}


def all_classes(connection: sqlite3.Connection) -> list[dict[str, Any]]:
    rows = connection.execute("SELECT id, name FROM classes ORDER BY name").fetchall()
    return [serialize_class(row, connection) for row in rows]


def extract_file_id(value: str) -> str | None:
    match = re.search(r"/drive/([A-Za-z0-9_-]+)", value or "")
    if not match:
        match = re.search(r"[?&]id=([A-Za-z0-9_-]+)", value or "")
    file_id = match.group(1) if match else (value if FILE_ID_RE.fullmatch(value or "") else None)
    return file_id if file_id and FILE_ID_RE.fullmatch(file_id) else None


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/state")
def state():
    with db() as connection:
        rows = connection.execute("SELECT * FROM submissions ORDER BY timestamp DESC, id DESC").fetchall()
        submissions = canonicalize_submission_names(rows)
        return jsonify({"classes": all_classes(connection), "submissions": submissions})


@app.post("/api/import")
def import_workbook():
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return jsonify(error="Choose an .xlsx workbook first."), 400
    if not uploaded.filename.lower().endswith(".xlsx"):
        return jsonify(error="Only .xlsx workbooks are supported."), 400

    try:
        records = records_from_xlsx(uploaded.stream)
    except Exception as exc:
        app.logger.exception("Workbook import failed")
        return jsonify(error=f"Could not read that workbook: {exc}"), 400

    replace_submissions(records)
    return jsonify(message=f"Loaded {len(records)} submissions.", count=len(records))


@app.post("/api/sync-master")
def sync_master_sheet():
    try:
        response = requests.get(
            app.config["MASTER_SHEET_EXPORT_URL"],
            timeout=(5, 30),
            stream=True,
        )
        response.raise_for_status()
        data = bytearray()
        for chunk in response.iter_content(64 * 1024):
            data.extend(chunk)
            if len(data) > MAX_UPLOAD_BYTES:
                return jsonify(error="The master Google Sheet is larger than 20 MB."), 413
        records = records_from_csv(bytes(data))
    except requests.RequestException:
        app.logger.exception("Master Google Sheet download failed")
        return jsonify(error="Could not download the latest submissions. Using the last saved copy."), 502
    except (csv.Error, UnicodeDecodeError, ValueError) as exc:
        app.logger.exception("Master Google Sheet import failed")
        return jsonify(error=f"Could not read the master Google Sheet: {exc}"), 502

    replace_submissions(records)
    synced_at = datetime.now().astimezone().isoformat(timespec="seconds")
    return jsonify(
        message=f"Updated {len(records)} submissions from the master Google Sheet.",
        count=len(records),
        synced_at=synced_at,
    )


@app.post("/api/classes")
def create_class():
    name = clean((request.get_json(silent=True) or {}).get("name"))
    if not name:
        return jsonify(error="Class name is required."), 400
    try:
        with db() as connection:
            cursor = connection.execute("INSERT INTO classes(name) VALUES (?)", (name,))
            row = connection.execute("SELECT id, name FROM classes WHERE id = ?", (cursor.lastrowid,)).fetchone()
            return jsonify(serialize_class(row, connection)), 201
    except sqlite3.IntegrityError:
        return jsonify(error="A class with that name already exists."), 409


@app.patch("/api/classes/<int:class_id>")
def rename_class(class_id: int):
    name = clean((request.get_json(silent=True) or {}).get("name"))
    if not name:
        return jsonify(error="Class name is required."), 400
    try:
        with db() as connection:
            cursor = connection.execute("UPDATE classes SET name = ? WHERE id = ?", (name, class_id))
            if not cursor.rowcount:
                return jsonify(error="Class not found."), 404
            row = connection.execute("SELECT id, name FROM classes WHERE id = ?", (class_id,)).fetchone()
            return jsonify(serialize_class(row, connection))
    except sqlite3.IntegrityError:
        return jsonify(error="A class with that name already exists."), 409


@app.delete("/api/classes/<int:class_id>")
def delete_class(class_id: int):
    with db() as connection:
        cursor = connection.execute("DELETE FROM classes WHERE id = ?", (class_id,))
        if not cursor.rowcount:
            return jsonify(error="Class not found."), 404
    return ("", 204)


@app.post("/api/classes/<int:class_id>/students")
def add_student(class_id: int):
    name = display_student_name((request.get_json(silent=True) or {}).get("name"))
    if not name:
        return jsonify(error="Student name is required."), 400
    try:
        with db() as connection:
            exists = connection.execute("SELECT 1 FROM classes WHERE id = ?", (class_id,)).fetchone()
            if not exists:
                return jsonify(error="Class not found."), 404
            position = connection.execute(
                "SELECT COALESCE(MAX(position), -1) + 1 FROM students WHERE class_id = ?", (class_id,)
            ).fetchone()[0]
            cursor = connection.execute(
                "INSERT INTO students(class_id, name, position) VALUES (?, ?, ?)", (class_id, name, position)
            )
            return jsonify(id=cursor.lastrowid, name=name, position=position), 201
    except sqlite3.IntegrityError:
        return jsonify(error="That student is already in this class."), 409


@app.patch("/api/students/<int:student_id>")
def rename_student(student_id: int):
    name = display_student_name((request.get_json(silent=True) or {}).get("name"))
    if not name:
        return jsonify(error="Student name is required."), 400
    try:
        with db() as connection:
            cursor = connection.execute("UPDATE students SET name = ? WHERE id = ?", (name, student_id))
            if not cursor.rowcount:
                return jsonify(error="Student not found."), 404
        return jsonify(id=student_id, name=name)
    except sqlite3.IntegrityError:
        return jsonify(error="That student is already in this class."), 409


@app.delete("/api/students/<int:student_id>")
def delete_student(student_id: int):
    with db() as connection:
        cursor = connection.execute("DELETE FROM students WHERE id = ?", (student_id,))
        if not cursor.rowcount:
            return jsonify(error="Student not found."), 404
    return ("", 204)


@app.get("/api/notebook/<file_id>")
def notebook(file_id: str):
    if not FILE_ID_RE.fullmatch(file_id):
        return jsonify(error="Invalid Google Drive file ID."), 400
    try:
        response = requests.get(
            "https://drive.google.com/uc", params={"id": file_id, "export": "download"}, timeout=20, stream=True
        )
        response.raise_for_status()
        data = bytearray()
        for chunk in response.iter_content(64 * 1024):
            data.extend(chunk)
            if len(data) > MAX_NOTEBOOK_BYTES:
                return jsonify(error="Notebook is larger than 15 MB."), 413
        parsed = json.loads(data.decode("utf-8"))
        if not isinstance(parsed.get("cells"), list):
            raise ValueError("File is not a Jupyter notebook")
        return jsonify(parsed)
    except (requests.RequestException, UnicodeDecodeError, json.JSONDecodeError, ValueError) as exc:
        return jsonify(error=f"Could not download the notebook. Check that its Drive link is shared: {exc}"), 502


@app.post("/api/resolve-notebook")
def resolve_notebook():
    file_id = extract_file_id(clean((request.get_json(silent=True) or {}).get("url")))
    if not file_id:
        return jsonify(error="This submission does not contain a valid Colab URL."), 400
    return jsonify(file_id=file_id)


@app.post("/api/execute-python")
def execute_python():
    payload = request.get_json(silent=True) or {}
    code = payload.get("code")
    if not isinstance(code, str) or not code.strip():
        return jsonify(error="Code is required.", output="", success=False), 400
    if len(code) > 100_000:
        return jsonify(error="Code is larger than 100 KB.", output="", success=False), 413
    with tempfile.TemporaryDirectory(prefix="sheet-explorer-") as directory:
        code_path = Path(directory) / "student_code.py"
        code_path.write_text(code, encoding="utf-8")
        try:
            completed = subprocess.run(
                [sys.executable, str(ROOT / "execution_runner.py"), str(code_path)],
                capture_output=True, text=True, timeout=30, cwd=directory,
            )
            result = json.loads(completed.stdout)
            return jsonify(result), (200 if result.get("success") else 422)
        except subprocess.TimeoutExpired:
            return jsonify(success=False, output="", error="Execution stopped after 30 seconds."), 408
        except (json.JSONDecodeError, OSError) as exc:
            return jsonify(success=False, output="", error=f"Execution failed: {exc}"), 500


@app.post("/api/execute-notebook")
def execute_notebook():
    notebook = (request.get_json(silent=True) or {}).get("notebook")
    if not isinstance(notebook, dict) or not isinstance(notebook.get("cells"), list):
        return jsonify(error="A valid notebook is required.", success=False), 400
    if len(notebook["cells"]) > 500:
        return jsonify(error="The notebook contains more than 500 cells.", success=False), 413
    source_size = sum(
        len(clean(cell.get("source")))
        for cell in notebook["cells"]
        if isinstance(cell, dict)
    )
    if source_size > 500_000:
        return jsonify(error="The notebook contains more than 500 KB of source code.", success=False), 413

    with tempfile.TemporaryDirectory(prefix="sheet-explorer-notebook-") as directory:
        notebook_path = Path(directory) / "submission.ipynb"
        notebook_path.write_text(json.dumps(notebook), encoding="utf-8")
        try:
            completed = subprocess.run(
                [sys.executable, str(ROOT / "notebook_runner.py"), str(notebook_path)],
                capture_output=True,
                text=True,
                timeout=180,
                cwd=directory,
            )
            if completed.returncode != 0:
                app.logger.error("Notebook runner failed: %s", completed.stderr[-2000:])
                return jsonify(error="The notebook kernel could not start.", success=False), 500
            return jsonify(json.loads(completed.stdout))
        except subprocess.TimeoutExpired:
            return jsonify(error="Run all stopped after three minutes.", success=False), 408
        except (json.JSONDecodeError, OSError) as exc:
            app.logger.exception("Notebook execution failed")
            return jsonify(error=f"Notebook execution failed: {exc}", success=False), 500


@app.errorhandler(413)
def too_large(_error):
    return jsonify(error="Upload is larger than 20 MB."), 413


init_db()

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "8000")), debug=True)
